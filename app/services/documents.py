import io
from pathlib import Path

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
ALLOWED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}


def extract_text(filename: str, content: bytes) -> str:
    if not content:
        raise ValueError("文件是空的")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("文件超过 10MB")

    ext = Path(filename or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError("仅支持 .txt、.md、.docx、.pdf")

    if ext == ".docx":
        text = _extract_docx(content)
    elif ext == ".pdf":
        text = _extract_pdf(content)
    else:
        text = _decode_text(content)

    text = text.strip()
    if not text:
        raise ValueError("未能从文件中提取文本")
    return text


def _decode_text(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("gbk", errors="ignore")


def _extract_docx(content: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(content))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n".join(page for page in pages if page)
