import io

import pytest
from docx import Document
from pypdf import PdfWriter

from app.services.documents import extract_text


def test_extract_utf8_and_gbk():
    assert extract_text("a.txt", "合同金额".encode("utf-8")) == "合同金额"
    assert "合同" in extract_text("a.txt", "合同".encode("gbk"))


def test_extract_docx_paragraph_and_table():
    document = Document()
    document.add_paragraph("采购合同正文")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "预付款"
    table.rows[0].cells[1].text = "30%"
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_text("合同.docx", buffer.getvalue())
    assert "采购合同正文" in text
    assert "预付款" in text
    assert "30%" in text


def test_extract_pdf_without_text_is_rejected():
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(ValueError, match="未能从文件中提取文本"):
        extract_text("blank.pdf", buffer.getvalue())


def test_reject_unknown_extension_and_oversize():
    with pytest.raises(ValueError, match="仅支持"):
        extract_text("scan.png", b"not-a-contract")
    with pytest.raises(ValueError, match="10MB"):
        extract_text("big.txt", b"a" * (10 * 1024 * 1024 + 1))
